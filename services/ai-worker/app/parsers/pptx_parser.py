"""Parser de PPTX con localización de diapositiva (E4-5).

Es el formato más importante del proyecto: el material que TCL China envía llega
como PowerPoint, y de él salen los claims que después hay que respaldar.

Se extrae, por diapositiva:
  - texto de todas las formas, incluidas las anidadas en grupos;
  - contenido de tablas, celda por celda — ahí viven las fichas técnicas;
  - notas del presentador, que a menudo contienen el detalle que falta.
"""

from __future__ import annotations

from typing import Any, Iterator, List

from .base import Fragment, ParseError, normalizar


def _formas(shapes: Any) -> Iterator[Any]:
    """Recorre las formas incluyendo las que están dentro de grupos."""
    for shape in shapes:
        if shape.shape_type == 6 and hasattr(shape, "shapes"):  # MSO_SHAPE_TYPE.GROUP
            yield from _formas(shape.shapes)
        else:
            yield shape


def _texto_de_tabla(shape: Any) -> str:
    filas = []
    for fila in shape.table.rows:
        celdas = [" ".join(celda.text.split()) for celda in fila.cells]
        if any(celdas):
            filas.append(" | ".join(celdas))
    return "\n".join(filas)


def _texto_de_diapositiva(slide: Any) -> tuple[str, str]:
    """Devuelve (cuerpo, título) de una diapositiva."""
    partes: List[str] = []
    titulo = ""

    for shape in _formas(slide.shapes):
        if shape.has_table:
            tabla = _texto_de_tabla(shape)
            if tabla:
                partes.append(tabla)
            continue

        if not shape.has_text_frame:
            continue

        texto = shape.text_frame.text.strip()
        if not texto:
            continue

        # El marcador de posición de título identifica el encabezado de la
        # diapositiva; sirve como `section` del fragmento.
        if not titulo and shape.is_placeholder and shape.placeholder_format.idx == 0:
            titulo = " ".join(texto.split())

        partes.append(texto)

    return "\n".join(partes), titulo


def _notas(slide: Any) -> str:
    if not slide.has_notes_slide:
        return ""
    frame = slide.notes_slide.notes_text_frame
    return frame.text.strip() if frame is not None else ""


def parse(path: str) -> List[Fragment]:
    try:
        from pptx import Presentation
    except ImportError as exc:  # pragma: no cover
        raise ParseError("Falta la librería de PowerPoint en el worker.") from exc

    try:
        presentacion = Presentation(path)
    except Exception as exc:
        raise ParseError(
            "No se pudo abrir el archivo de PowerPoint. Puede estar dañado o protegido."
        ) from exc

    fragmentos: List[Fragment] = []
    ordinal = 0

    for numero, slide in enumerate(presentacion.slides, start=1):
        try:
            cuerpo, titulo = _texto_de_diapositiva(slide)
            notas = _notas(slide)
        except Exception as exc:
            raise ParseError(
                f"No se pudo leer la diapositiva {numero}: el archivo puede estar dañado."
            ) from exc

        if cuerpo.strip():
            fragmentos.append(
                Fragment(
                    ordinal=ordinal,
                    text=normalizar(cuerpo),
                    slide=numero,
                    section=titulo or None,
                    metadata={"origen": "diapositiva"},
                )
            )
            ordinal += 1

        # Las notas van como fragmento aparte: son contexto del presentador, no
        # contenido publicado, y conviene poder distinguirlos al citar.
        if notas.strip():
            fragmentos.append(
                Fragment(
                    ordinal=ordinal,
                    text=normalizar(notas),
                    slide=numero,
                    section=titulo or None,
                    metadata={"origen": "notas"},
                )
            )
            ordinal += 1

    if not fragmentos:
        raise ParseError(
            "El archivo no contiene texto extraíble. "
            "Si las diapositivas son imágenes, se requiere OCR."
        )

    return fragmentos
