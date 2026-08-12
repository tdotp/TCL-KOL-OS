"""Parsers con localización de página y diapositiva (E4-5, E4-6).

La propiedad que se prueba una y otra vez: **todo fragmento sabe de dónde salió**.
Sin eso no hay evidencia verificable, y sin evidencia el umbral de §15
("claims con evidencia: 100%") es inalcanzable.
"""

from __future__ import annotations

from pathlib import Path

import pytest

from app import parsers
from app.parsers.base import Fragment, ParseError


# ─────────────────────────── Contrato de Fragment ─────────────────────────────


def test_un_fragmento_sin_ubicacion_es_invalido() -> None:
    with pytest.raises(ValueError, match="evidencia"):
        Fragment(ordinal=0, text="Panel de 144Hz")


def test_un_fragmento_vacio_es_invalido() -> None:
    with pytest.raises(ValueError, match="vacío"):
        Fragment(ordinal=0, text="   ", slide=1)


def test_un_fragmento_con_diapositiva_es_valido() -> None:
    f = Fragment(ordinal=0, text="Panel de 144Hz", slide=11)
    assert f.slide == 11 and f.page is None


# ─────────────────────────── PPTX ─────────────────────────────────────────────


def crear_pptx(ruta: Path) -> Path:
    from pptx import Presentation
    from pptx.util import Inches

    pres = Presentation()

    # Diapositiva 1: título y contenido.
    s1 = pres.slides.add_slide(pres.slide_layouts[1])
    s1.shapes.title.text = "Ficha técnica C755"
    s1.placeholders[1].text = "Panel nativo de 144Hz\nMini LED"

    # Diapositiva 2: tabla con especificaciones.
    s2 = pres.slides.add_slide(pres.slide_layouts[5])
    s2.shapes.title.text = "Comparativa"
    tabla = s2.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(6), Inches(1)).table
    tabla.cell(0, 0).text = "Brillo"
    tabla.cell(0, 1).text = "1300 nits"
    tabla.cell(1, 0).text = "Almacenamiento"
    tabla.cell(1, 1).text = "512GB"

    # Diapositiva 3: solo notas del presentador.
    s3 = pres.slides.add_slide(pres.slide_layouts[6])
    s3.notes_slide.notes_text_frame.text = "Confirmar el precio con TCL Colombia."

    pres.save(str(ruta))
    return ruta


def test_pptx_numera_las_diapositivas_desde_uno(tmp_path: Path) -> None:
    ruta = crear_pptx(tmp_path / "campana.pptx")
    frags = parsers.parse(str(ruta), parsers.PPTX)

    assert all(f.slide is not None for f in frags)
    assert min(f.slide for f in frags) == 1
    assert max(f.slide for f in frags) == 3


def test_pptx_extrae_el_contenido_de_las_tablas(tmp_path: Path) -> None:
    """Las fichas técnicas viven en tablas: perderlas sería perder los claims."""
    ruta = crear_pptx(tmp_path / "campana.pptx")
    frags = parsers.parse(str(ruta), parsers.PPTX)

    diapositiva2 = [f for f in frags if f.slide == 2]
    texto = " ".join(f.text for f in diapositiva2)
    assert "1300 nits" in texto
    assert "512GB" in texto


def test_pptx_usa_el_titulo_como_seccion(tmp_path: Path) -> None:
    ruta = crear_pptx(tmp_path / "campana.pptx")
    frags = parsers.parse(str(ruta), parsers.PPTX)

    primera = next(f for f in frags if f.slide == 1)
    assert primera.section == "Ficha técnica C755"


def test_pptx_separa_las_notas_del_contenido(tmp_path: Path) -> None:
    """Las notas son contexto del presentador, no contenido publicado."""
    ruta = crear_pptx(tmp_path / "campana.pptx")
    frags = parsers.parse(str(ruta), parsers.PPTX)

    notas = [f for f in frags if f.metadata.get("origen") == "notas"]
    assert len(notas) == 1
    assert notas[0].slide == 3
    assert "Confirmar el precio" in notas[0].text


def test_pptx_ordinales_sin_huecos_ni_repetidos(tmp_path: Path) -> None:
    """La idempotencia de E4-10 se apoya en `UNIQUE(source_id, ordinal)`."""
    ruta = crear_pptx(tmp_path / "campana.pptx")
    frags = parsers.parse(str(ruta), parsers.PPTX)

    ordinales = [f.ordinal for f in frags]
    assert ordinales == list(range(len(frags)))


def test_pptx_sin_texto_falla_con_mensaje_accionable(tmp_path: Path) -> None:
    from pptx import Presentation

    pres = Presentation()
    pres.slides.add_slide(pres.slide_layouts[6])
    ruta = tmp_path / "vacia.pptx"
    pres.save(str(ruta))

    with pytest.raises(ParseError, match="OCR"):
        parsers.parse(str(ruta), parsers.PPTX)


# ─────────────────────────── DOCX ─────────────────────────────────────────────


def crear_docx(ruta: Path, con_salto: bool = True) -> Path:
    import docx

    d = docx.Document()
    d.add_heading("Contexto", level=1)
    d.add_paragraph("El mercado está dominado por Samsung y Xiaomi.")

    if con_salto:
        d.add_page_break()

    d.add_heading("Mandatorios", level=1)
    d.add_paragraph("Aclarar que los 144Hz son nativos del panel.")

    tabla = d.add_table(rows=2, cols=2)
    tabla.cell(0, 0).text = "Modelo"
    tabla.cell(0, 1).text = "Precio"
    tabla.cell(1, 0).text = "P80 Ultra"
    tabla.cell(1, 1).text = "$2.999.900"

    d.save(str(ruta))
    return ruta


def test_docx_cuenta_las_paginas_por_saltos_explicitos(tmp_path: Path) -> None:
    ruta = crear_docx(tmp_path / "brief.docx", con_salto=True)
    frags = parsers.parse(str(ruta), parsers.DOCX)

    paginas = {f.page for f in frags}
    assert paginas == {1, 2}, "el salto de página explícito debe incrementar el número"


def test_docx_sin_saltos_queda_todo_en_la_pagina_uno(tmp_path: Path) -> None:
    """Un DOCX no tiene páginas hasta renderizarse: se declara la aproximación."""
    ruta = crear_docx(tmp_path / "brief.docx", con_salto=False)
    frags = parsers.parse(str(ruta), parsers.DOCX)

    assert {f.page for f in frags} == {1}
    assert all(f.metadata.get("paginacion") == "saltos_explicitos" for f in frags)


def test_docx_atribuye_cada_bloque_a_su_encabezado(tmp_path: Path) -> None:
    ruta = crear_docx(tmp_path / "brief.docx")
    frags = parsers.parse(str(ruta), parsers.DOCX)

    secciones = {f.section for f in frags if f.section is not None}
    assert "Contexto" in secciones
    assert "Mandatorios" in secciones


def test_docx_extrae_las_tablas_por_separado(tmp_path: Path) -> None:
    ruta = crear_docx(tmp_path / "brief.docx")
    frags = parsers.parse(str(ruta), parsers.DOCX)

    tablas = [f for f in frags if f.metadata.get("origen") == "tabla"]
    assert len(tablas) == 1
    assert "P80 Ultra" in tablas[0].text
    assert "$2.999.900" in tablas[0].text


# ─────────────────────────── Registro ─────────────────────────────────────────


def test_un_tipo_no_soportado_falla_sin_perder_el_original(tmp_path: Path) -> None:
    ruta = tmp_path / "hoja.xlsx"
    ruta.write_bytes(b"PK\x03\x04")

    with pytest.raises(ParseError, match="conserva el original"):
        parsers.parse(str(ruta), "application/vnd.ms-excel")


def test_soporta_declara_los_formatos_con_extraccion() -> None:
    assert parsers.soporta(parsers.PPTX)
    assert parsers.soporta(parsers.PDF)
    assert parsers.soporta(parsers.DOCX)
    # Las imágenes se aceptan al cargar, pero su extracción exige OCR (E4-7).
    assert not parsers.soporta("image/png")
